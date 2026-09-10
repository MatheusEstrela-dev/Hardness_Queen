import importlib.util
import json
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from ingestao.extracao_texto import (
    TAMANHO_MAXIMO_DO_CHUNK,
    extrair,
    extrair_docx,
    extrair_pdf,
    titulo_da_secao,
)

_SCRIPT_CLI = Path(__file__).resolve().parents[1] / "scripts" / "03_extrair_texto.py"


def _carregar_cli():
    especificacao = importlib.util.spec_from_file_location("cli_extrair_texto", _SCRIPT_CLI)
    modulo = importlib.util.module_from_spec(especificacao)
    especificacao.loader.exec_module(modulo)
    return modulo


def _pdf_com_texto(caminho: Path) -> Path:
    documento = pymupdf.open()
    pagina = documento.new_page()
    pagina.insert_text((72, 100), "4.2 Caracterizacao do solo gnaisse", fontsize=16)
    pagina.insert_text((72, 140), "A saturacao ocorre a partir de 100mm em 72h.", fontsize=11)
    documento.save(caminho)
    documento.close()
    return caminho


def _pdf_sem_camada_texto(caminho: Path) -> Path:
    documento = pymupdf.open()
    documento.new_page()
    documento.save(caminho)
    documento.close()
    return caminho


def _docx_com_secao(caminho: Path) -> Path:
    documento = Document()
    documento.add_heading("1 Limiares hidrologicos", level=1)
    documento.add_paragraph("A cota de transbordamento do Rio Arrudas e 3.8m.")
    documento.save(caminho)
    return caminho


def test_pdf_produz_chunk_com_numero_de_pagina(tmp_path: Path):
    chunks, estado = extrair_pdf(_pdf_com_texto(tmp_path / "laudo.pdf"))

    assert estado == "texto_nativo"
    assert len(chunks) == 1
    assert chunks[0].pagina == 1


def test_pdf_preserva_o_corpo_da_pagina_no_chunk(tmp_path: Path):
    chunks, _ = extrair_pdf(_pdf_com_texto(tmp_path / "laudo.pdf"))

    assert "gnaisse" in chunks[0].texto
    assert "100mm" in chunks[0].texto


def test_pdf_com_tabela_preserva_as_colunas_em_markdown(tmp_path: Path):
    caminho = tmp_path / "tabela.pdf"
    documento = pymupdf.open()
    pagina = documento.new_page()
    html = """
    <table border="1">
      <tr><th>Solo</th><th>Limiar 72h (mm)</th></tr>
      <tr><td>gnaisse</td><td>100</td></tr>
      <tr><td>argiloso</td><td>80</td></tr>
    </table>
    """
    pagina.insert_htmlbox(pymupdf.Rect(50, 50, 500, 300), html)
    documento.save(caminho)
    documento.close()

    chunks, _ = extrair_pdf(caminho)

    texto = chunks[0].texto
    assert "gnaisse" in texto and "argiloso" in texto
    assert "100" in texto and "80" in texto
    # a associacao solo -> valor precisa sobreviver: gnaisse e 100 na mesma linha
    linha_do_gnaisse = next(l for l in texto.splitlines() if "gnaisse" in l)
    assert "100" in linha_do_gnaisse


def _pdf_multipagina_com_secoes(caminho: Path) -> Path:
    documento = pymupdf.open()

    pagina1 = documento.new_page()
    pagina1.insert_text((72, 100), "4.2 Caracterizacao do solo gnaisse", fontsize=16)
    pagina1.insert_text((72, 140), "Corpo da secao 4.2 na pagina 1.", fontsize=11)

    pagina2 = documento.new_page()
    pagina2.insert_text((72, 100), "A saturacao ocorre a partir de 100mm em 72h.", fontsize=11)

    pagina3 = documento.new_page()
    pagina3.insert_text((72, 100), "5.1 Limiares de vazao", fontsize=16)
    pagina3.insert_text((72, 140), "Corpo da secao 5.1 na pagina 3.", fontsize=11)

    documento.save(caminho)
    documento.close()
    return caminho


def test_pdf_propaga_secao_para_pagina_seguinte_sem_cabecalho(tmp_path: Path):
    chunks, _ = extrair_pdf(_pdf_multipagina_com_secoes(tmp_path / "multipagina.pdf"))

    assert len(chunks) == 3
    # pagina 2 nao tem cabecalho proprio: herda a secao da pagina 1
    assert chunks[1].secao == "4.2 Caracterizacao do solo gnaisse"


def test_pdf_novo_cabecalho_substitui_a_secao_herdada(tmp_path: Path):
    chunks, _ = extrair_pdf(_pdf_multipagina_com_secoes(tmp_path / "multipagina.pdf"))

    # pagina 3 tem cabecalho proprio: nao deve vazar a secao da pagina 1
    assert chunks[2].secao == "5.1 Limiares de vazao"


def test_pdf_sem_camada_de_texto_nao_produz_zero_em_silencio(tmp_path: Path):
    chunks, estado = extrair_pdf(_pdf_sem_camada_texto(tmp_path / "digitalizado.pdf"))

    assert estado in {"sem_camada_texto", "ocr_aplicado"}
    if estado == "sem_camada_texto":
        assert chunks == []


def _docx_com_multiplos_headings(caminho: Path) -> Path:
    documento = Document()
    documento.add_heading("1 Limiares hidrologicos", level=1)
    documento.add_paragraph("Corpo da secao 1.")
    documento.add_heading("2 Limiares geologicos", level=1)
    documento.add_paragraph("Corpo da secao 2.")
    documento.save(caminho)
    return caminho


def _paragrafo_em_negrito(documento: Document, texto: str):
    paragrafo = documento.add_paragraph()
    corrida = paragrafo.add_run(texto)
    corrida.bold = True
    return paragrafo


def _docx_sem_estilo_com_negrito_curto(caminho: Path) -> Path:
    # documento real (Instrucao_Normativa_COMPLETA_Alertas_MG.docx) marca
    # titulos so com negrito/maiuscula em paragrafos "Normal", sem usar
    # estilo Heading -- e o padrao que este fixture reproduz.
    documento = Document()
    _paragrafo_em_negrito(documento, "4.2 CARACTERIZACAO DO SOLO")
    documento.add_paragraph("A cota de transbordamento do Rio Arrudas e 3.8m.")
    documento.add_paragraph("A saturacao ocorre a partir de 100mm em 72h.")
    documento.save(caminho)
    return caminho


def _docx_corpo_longo_sem_cabecalho(caminho: Path) -> Path:
    documento = Document()
    for indice in range(40):
        documento.add_paragraph(
            f"Paragrafo {indice} de corpo continuo sem nenhum marcador de secao, "
            "escrito para simular prosa real de um laudo tecnico sem estilo de titulo."
        )
    documento.save(caminho)
    return caminho


def _docx_secao_com_corpo_que_estoura_o_limite(caminho: Path) -> Path:
    documento = Document()
    _paragrafo_em_negrito(documento, "6.2 LIMIARES DE CHUVA")
    for indice in range(40):
        documento.add_paragraph(
            f"Paragrafo {indice} do corpo da secao 6.2, prosa longa o bastante para "
            "estourar o limite de tamanho do chunk depois de varias repeticoes."
        )
    documento.save(caminho)
    return caminho


def _docx_negrito_no_meio_de_paragrafo_longo(caminho: Path) -> Path:
    documento = Document()
    documento.add_heading("1 Introducao", level=1)
    paragrafo = documento.add_paragraph()
    paragrafo.add_run(
        "Texto normal antes da parte em destaque, com bastante contexto para "
        "deixar o paragrafo longo de verdade. "
    )
    corrida_negrito = paragrafo.add_run("Trecho em negrito no meio da frase")
    corrida_negrito.bold = True
    paragrafo.add_run(
        ", seguido de mais texto normal para completar a prosa e ultrapassar "
        "o tamanho tipico de um cabecalho de verdade."
    )
    documento.save(caminho)
    return caminho


def test_docx_com_varios_headings_de_estilo_produz_um_chunk_por_secao(tmp_path: Path):
    # regressao: documento que usa o estilo Heading continua se dividindo por
    # ele, sem a heuristica de negrito/maiuscula interferir.
    chunks, _ = extrair_docx(_docx_com_multiplos_headings(tmp_path / "laudo.docx"))

    assert len(chunks) == 2
    assert chunks[0].secao == "1 Limiares hidrologicos"
    assert chunks[1].secao == "2 Limiares geologicos"


def test_docx_sem_estilo_de_heading_usa_negrito_curto_como_secao(tmp_path: Path):
    chunks, _ = extrair_docx(_docx_sem_estilo_com_negrito_curto(tmp_path / "laudo.docx"))

    assert len(chunks) == 1
    assert chunks[0].secao == "4.2 CARACTERIZACAO DO SOLO"
    assert "3.8m" in chunks[0].texto
    assert "100mm" in chunks[0].texto


def test_docx_sem_cabecalho_detectavel_e_dividido_pelo_limite_de_tamanho(tmp_path: Path):
    chunks, _ = extrair_docx(_docx_corpo_longo_sem_cabecalho(tmp_path / "laudo.docx"))

    assert len(chunks) > 1
    assert all(len(chunk.texto) <= TAMANHO_MAXIMO_DO_CHUNK for chunk in chunks)


def test_docx_chunk_de_continuacao_carrega_a_secao_do_bloco(tmp_path: Path):
    chunks, _ = extrair_docx(_docx_secao_com_corpo_que_estoura_o_limite(tmp_path / "laudo.docx"))

    assert len(chunks) > 1
    assert all(chunk.secao == "6.2 LIMIARES DE CHUVA" for chunk in chunks)


def test_docx_negrito_no_meio_de_frase_longa_nao_vira_secao(tmp_path: Path):
    # guarda contra falso positivo: uma frase de prosa com um trecho em
    # negrito no meio (nao o paragrafo inteiro) nao pode virar cabecalho.
    chunks, _ = extrair_docx(_docx_negrito_no_meio_de_paragrafo_longo(tmp_path / "laudo.docx"))

    assert len(chunks) == 1
    assert chunks[0].secao == "1 Introducao"


def test_docx_produz_pagina_nula_e_secao_preenchida(tmp_path: Path):
    chunks, estado = extrair_docx(_docx_com_secao(tmp_path / "laudo.docx"))

    assert estado == "texto_nativo"
    assert chunks[0].pagina is None
    assert chunks[0].secao == "1 Limiares hidrologicos"


def _docx_paragrafo_tabela_paragrafo(caminho: Path) -> Path:
    documento = Document()
    documento.add_paragraph("Paragrafo antes da tabela.")
    tabela = documento.add_table(rows=2, cols=2)
    tabela.rows[0].cells[0].text = "Intensidade"
    tabela.rows[0].cells[1].text = "Taxa de precipitacao"
    tabela.rows[1].cells[0].text = "Forte"
    tabela.rows[1].cells[1].text = "30 mm/h a 70 mm/h"
    documento.add_paragraph("Paragrafo depois da tabela.")
    documento.save(caminho)
    return caminho


def _docx_tabela_perto_do_limite(caminho: Path) -> Path:
    # preenchimento perto do teto (2800): a tabela so cabe se o bloco
    # anterior for fechado primeiro -- e a condicao que exercita "tabela
    # nunca dividida entre chunks".
    documento = Document()
    texto_grande = ("Preenchimento de prosa para chegar perto do teto do chunk. " * 60)[:2760]
    documento.add_paragraph(texto_grande)
    tabela = documento.add_table(rows=2, cols=2)
    tabela.rows[0].cells[0].text = "PRIMEIRA_CELULA"
    tabela.rows[0].cells[1].text = "cabecalho"
    tabela.rows[1].cells[0].text = "corpo"
    tabela.rows[1].cells[1].text = "ULTIMA_CELULA"
    documento.save(caminho)
    return caminho


def _docx_tabela_vazia_entre_paragrafos(caminho: Path) -> Path:
    documento = Document()
    documento.add_paragraph("Antes da tabela vazia.")
    documento.add_table(rows=0, cols=2)
    documento.add_paragraph("Depois da tabela vazia.")
    documento.save(caminho)
    return caminho


def _docx_tabela_com_celulas_vazias(caminho: Path) -> Path:
    documento = Document()
    documento.add_paragraph("Antes.")
    tabela = documento.add_table(rows=2, cols=2)
    tabela.rows[0].cells[0].text = "Cabecalho"
    # as demais celulas ficam vazias de proposito
    documento.add_paragraph("Depois.")
    documento.save(caminho)
    return caminho


def test_docx_tabela_aparece_entre_os_paragrafos_na_ordem_do_documento(tmp_path: Path):
    chunks, _ = extrair_docx(_docx_paragrafo_tabela_paragrafo(tmp_path / "laudo.docx"))

    assert len(chunks) == 1
    texto = chunks[0].texto
    posicao_antes = texto.index("Paragrafo antes da tabela.")
    posicao_tabela = texto.index("Forte")
    posicao_depois = texto.index("Paragrafo depois da tabela.")
    assert posicao_antes < posicao_tabela < posicao_depois


def test_docx_tabela_mantem_linha_e_valor_na_mesma_linha(tmp_path: Path):
    chunks, _ = extrair_docx(_docx_paragrafo_tabela_paragrafo(tmp_path / "laudo.docx"))

    texto = chunks[0].texto
    linha_forte = next(linha for linha in texto.splitlines() if "Forte" in linha)
    assert "30 mm/h a 70 mm/h" in linha_forte


def test_docx_tabela_preserva_linha_de_cabecalho_markdown(tmp_path: Path):
    chunks, _ = extrair_docx(_docx_paragrafo_tabela_paragrafo(tmp_path / "laudo.docx"))

    texto = chunks[0].texto
    linhas = texto.splitlines()
    indice_cabecalho = next(i for i, linha in enumerate(linhas) if "Intensidade" in linha)
    linha_separadora = linhas[indice_cabecalho + 1]
    assert linha_separadora.strip("|")
    assert set(linha_separadora) <= {"-", "|"}


def test_docx_tabela_nunca_e_dividida_entre_chunks(tmp_path: Path):
    chunks, _ = extrair_docx(_docx_tabela_perto_do_limite(tmp_path / "laudo.docx"))

    assert len(chunks) >= 2
    chunk_com_tabela = next(chunk for chunk in chunks if "PRIMEIRA_CELULA" in chunk.texto)
    assert "ULTIMA_CELULA" in chunk_com_tabela.texto


def test_docx_tabela_vazia_nao_quebra_a_extracao(tmp_path: Path):
    chunks, _ = extrair_docx(_docx_tabela_vazia_entre_paragrafos(tmp_path / "laudo.docx"))

    assert len(chunks) == 1
    assert "Antes da tabela vazia." in chunks[0].texto
    assert "Depois da tabela vazia." in chunks[0].texto


def test_docx_tabela_com_celulas_vazias_nao_produz_linha_malformada(tmp_path: Path):
    chunks, _ = extrair_docx(_docx_tabela_com_celulas_vazias(tmp_path / "laudo.docx"))

    texto = chunks[0].texto
    linhas_da_tabela = [linha for linha in texto.splitlines() if linha.startswith("|")]
    assert linhas_da_tabela
    numero_de_colunas = linhas_da_tabela[0].count("|")
    assert all(linha.count("|") == numero_de_colunas for linha in linhas_da_tabela)


def test_docx_sem_tabela_mantem_a_mesma_divisao_de_chunks_de_antes(tmp_path: Path):
    # regressao: documento sem nenhuma tabela produz exatamente a mesma
    # divisao por secao de antes desta mudanca.
    chunks, _ = extrair_docx(_docx_com_multiplos_headings(tmp_path / "laudo.docx"))

    assert len(chunks) == 2
    assert chunks[0].secao == "1 Limiares hidrologicos"
    assert chunks[1].secao == "2 Limiares geologicos"


def test_chunk_id_e_preenchido(tmp_path: Path):
    chunks, _ = extrair_pdf(_pdf_com_texto(tmp_path / "laudo.pdf"))

    assert len(chunks[0].chunk_id) == 16


def test_extrair_despacha_por_extensao(tmp_path: Path):
    chunks_pdf, _ = extrair(_pdf_com_texto(tmp_path / "a.pdf"))
    chunks_docx, _ = extrair(_docx_com_secao(tmp_path / "a.docx"))

    assert chunks_pdf and chunks_docx


def test_extrair_rejeita_extensao_desconhecida(tmp_path: Path):
    arquivo = tmp_path / "planilha.xlsx"
    arquivo.write_text("x", encoding="utf-8")

    with pytest.raises(ValueError):
        extrair(arquivo)


def test_titulo_da_secao_pega_o_primeiro_cabecalho():
    assert titulo_da_secao("## 4.2 Solo gnaisse\n\ncorpo") == "4.2 Solo gnaisse"


def test_titulo_da_secao_sem_cabecalho_devolve_nulo():
    assert titulo_da_secao("apenas corpo de texto") is None


def test_cli_documento_invalido_nao_interrompe_o_lote(tmp_path: Path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    pasta_docs = tmp_path / "docs"
    pasta_docs.mkdir()
    (pasta_docs / "a_corrompido.pdf").write_bytes(b"nao sou um pdf valido")
    _pdf_com_texto(pasta_docs / "b_valido.pdf")

    modulo = _carregar_cli()
    resultado = modulo.main()

    saida = capsys.readouterr().out
    assert resultado == 0
    # documento invalido nao interrompe o lote: o proximo documento ainda e processado
    assert "[erro_extracao] a_corrompido.pdf" in saida
    assert "[texto_nativo] b_valido.pdf: 1 chunks" in saida
    assert "erro_extracao: 1" in saida
    assert "texto_nativo: 1" in saida

    registros = {
        registro["doc"]: registro
        for registro in (
            json.loads(linha)
            for linha in (tmp_path / "data" / "manifesto.jsonl").read_text(encoding="utf-8").splitlines()
        )
    }
    assert registros["a_corrompido.pdf"]["estado"] == "erro_extracao"
    assert registros["b_valido.pdf"]["estado"] == "texto_nativo"
